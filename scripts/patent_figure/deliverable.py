"""附图导出包的页面合成（F1，WP2）：规范版本 DXF → A4 页面栅格 → PDF / Word。

只服务 ``patent_figure.export``（单一导出服务）；读的是被钉住的版本目录里的 DXF，写的只在导出
包目录里。OCP 无关；栅格化用 ezdxf 的 matplotlib 后端（Agg，面向对象 API，不碰 pyplot 全局状态，
导出作业线程里安全），PDF 用 matplotlib 的 ``PdfPages``，Word 用可选依赖 python-docx。**出图链路
（DXF、N10 金样）一个字节都不受影响**：版式参数只在这里生效。

版式（任务书 F1，契约 §14.1）：

* A4（210×297 mm）竖放，页边 上 25 / 左 25 / 下 15 / 右 15 mm → 可用区 170×257 mm；
* 图区预设 ``standard`` = 170×257、``conservative`` = 165×245，居中放在可用区内；
* 每页一张图：取 DXF 全部图元的外廓（含图号、件号表，外扩 ``WINDOW_PAD_MM``）为取景窗，
  **不放大**——缩放比 = min(1, 图区宽/窗宽, 图区高/窗高)，所以字高只会不变或变小，由
  ``deliverable_check`` 逐页核验缩放后字高；
* 栅格 dpi 可配（默认 300，范围 72–600），灰度 8 位、白底黑线；
* 页序 = 图号序（图 1 … 图 n，即 plan.figures 的顺序；「图 N」已由标注步骤写进规范版本 DXF）。

精确图片参数以当期《电子申请文件格式要求说明》为准（``CNIPA_NOTE``，能力接口原样返回）。
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from patent_figure import layout as _layout

PAGE_MM = (210.0, 297.0)
MARGINS_MM = {"top": 25.0, "left": 25.0, "bottom": 15.0, "right": 15.0}
SHEET_PRESETS = {
    "standard": {"area_mm": (170.0, 257.0), "label": "标准版心 170×257 mm"},
    "conservative": {"area_mm": (165.0, 245.0), "label": "保守版心 165×245 mm"},
}
DEFAULT_PRESET = "standard"
DEFAULT_DPI = 300
DPI_MIN, DPI_MAX = 72, 600
#: Clearance added around the drawing's ink extents to form the page's capture window (mm).
WINDOW_PAD_MM = 2.0
#: 缩放后最小字高（mm）：A16 / 审查指南五部一章 5.2 的 3.5 mm，与出图闸门同一常数。
MIN_TEXT_MM = float(_layout.TEXT_FLOOR_MM)
#: 栅格上最小字高（像素）。汉字笔画可辨约需 16 px 字身；A11 要求缩印到 2/3 仍清晰 →
#: 16 / (2/3) = 24 px。300 dpi 下 3.5 mm ≈ 41 px（通过），72 dpi 下 ≈ 9.9 px（不通过）。
MIN_TEXT_PX = 24.0
CNIPA_NOTE = "精确图片参数以当期《电子申请文件格式要求说明》为准"
#: Gray level (0–255) below which a pixel counts as ink.
INK_LEVEL = 128


class DeliverableError(RuntimeError):
    pass


def spec(preset: str = DEFAULT_PRESET, dpi: int = DEFAULT_DPI) -> Dict[str, Any]:
    """Everything a page is built against and checked against (also the capabilities shape)."""
    if preset not in SHEET_PRESETS:
        raise DeliverableError("未知图区预设 %r（可选：%s）" % (preset, "、".join(SHEET_PRESETS)))
    dpi = int(dpi)
    if not (DPI_MIN <= dpi <= DPI_MAX):
        raise DeliverableError("dpi 必须在 %d–%d 之间，收到 %d" % (DPI_MIN, DPI_MAX, dpi))
    aw, ah = SHEET_PRESETS[preset]["area_mm"]
    return {"page_mm": list(PAGE_MM), "margins_mm": dict(MARGINS_MM), "preset": preset,
            "area_mm": [aw, ah], "dpi": dpi, "min_text_mm": MIN_TEXT_MM,
            "min_text_px": MIN_TEXT_PX, "window_pad_mm": WINDOW_PAD_MM, "note": CNIPA_NOTE}


def printable_mm() -> Tuple[float, float]:
    return (PAGE_MM[0] - MARGINS_MM["left"] - MARGINS_MM["right"],
            PAGE_MM[1] - MARGINS_MM["top"] - MARGINS_MM["bottom"])


# --------------------------------------------------------------------------- DXF facts


def read_doc(dxf: Path):
    import ezdxf
    return ezdxf.readfile(str(dxf))


def sheet_window(msp, area_mm: Optional[Sequence[float]] = None
                 ) -> Tuple[float, float, float, float]:
    """The capture window (DXF mm): ink extents of every entity plus a clearance of up to
    ``WINDOW_PAD_MM`` per side — reduced (never below 0) on an axis where the full clearance
    would push a drawing that FITS the drawing area over it, so a fitting drawing is always
    placed 1:1 and its text heights are untouched. Shared by the writer and the checker; derived
    from the SOURCE drawing and the chosen preset only, never from what the writer reports."""
    from ezdxf import bbox
    ext = bbox.extents(msp)
    if not ext.has_data:
        raise DeliverableError("DXF 里没有任何图元")
    x0, y0 = float(ext.extmin.x), float(ext.extmin.y)
    x1, y1 = float(ext.extmax.x), float(ext.extmax.y)
    px = py = WINDOW_PAD_MM
    if area_mm is not None:
        aw, ah = float(area_mm[0]), float(area_mm[1])
        if x1 - x0 <= aw:
            px = max(0.0, min(px, (aw - (x1 - x0)) / 2.0))
        if y1 - y0 <= ah:
            py = max(0.0, min(py, (ah - (y1 - y0)) / 2.0))
    return (x0 - px, y0 - py, x1 + px, y1 + py)


def text_heights(msp) -> List[Tuple[str, str, float]]:
    """(layer, text, height mm) of every TEXT / MTEXT in the drawing."""
    out = []
    for e in msp:
        kind = e.dxftype()
        if kind == "TEXT":
            out.append((str(e.dxf.layer), str(e.dxf.text), float(e.dxf.height)))
        elif kind == "MTEXT":
            out.append((str(e.dxf.layer), str(e.text), float(e.dxf.get("char_height", 0.0))))
    return out


def caption_of(msp) -> Optional[Dict[str, Any]]:
    """The figure-number caption (CAPTION layer TEXT): text, anchor and alignment."""
    from patent_figure import qa as _qa
    for e in msp:
        if e.dxftype() in ("TEXT", "MTEXT") and str(e.dxf.layer) == "CAPTION":
            payload = _qa._text_payload(e)
            if payload:
                text, h, anchor, hk, vk = payload
                return {"text": text.strip(), "height": h, "anchor": anchor, "h": hk, "v": vk}
    return None


def caption_box(cap: Dict[str, Any]) -> Tuple[float, float, float, float]:
    """A box (DXF mm) that certainly overlaps the caption glyphs: CJK-aware width estimate,
    then the central 80 % (so a generous estimate cannot reach neighbouring ink)."""
    h = float(cap["height"])
    tw = sum((0.62 if ord(c) < 0x2E80 else 1.0) for c in cap["text"]) * h
    ax, ay = cap["anchor"]
    x0 = ax - tw if cap["h"] == "R" else (ax - tw / 2.0 if cap["h"] == "C" else ax)
    yc = ay if cap["v"] == "MID" else (ay - h / 2.0 if cap["v"] == "TOP" else ay + h / 2.0)
    box = (x0, yc - h / 2.0, x0 + tw, yc + h / 2.0)
    cx, cy = (box[0] + box[2]) / 2.0, (box[1] + box[3]) / 2.0
    return (cx - 0.4 * (box[2] - box[0]), cy - 0.4 * (box[3] - box[1]),
            cx + 0.4 * (box[2] - box[0]), cy + 0.4 * (box[3] - box[1]))


def place(window: Sequence[float], sp: Dict[str, Any]) -> Dict[str, float]:
    """Page placement (mm, origin top-left) of a capture window under preset ``sp``."""
    ww, wh = window[2] - window[0], window[3] - window[1]
    aw, ah = sp["area_mm"]
    pw, ph = printable_mm()
    scale = min(1.0, aw / ww, ah / wh)
    iw, ih = ww * scale, wh * scale
    ax0 = MARGINS_MM["left"] + (pw - aw) / 2.0
    at0 = MARGINS_MM["top"] + (ph - ah) / 2.0
    return {"scale": scale, "w": iw, "h": ih,
            "x": ax0 + (aw - iw) / 2.0, "top": at0 + (ah - ih) / 2.0,
            "area_x": ax0, "area_top": at0}


# --------------------------------------------------------------------------- raster


class Page:
    """One rasterised page: the gray image + where it goes + the facts it was built from."""

    def __init__(self, figure: str, number: int, dxf: Path, gray, window, placement, dpi):
        self.figure, self.number, self.dxf = figure, number, Path(dxf)
        self.gray, self.window, self.placement, self.dpi = gray, window, placement, dpi

    def png_bytes(self) -> bytes:
        from PIL import Image
        buf = io.BytesIO()
        Image.fromarray(self.gray, "L").save(buf, "PNG", dpi=(self.dpi, self.dpi))
        return buf.getvalue()


def rasterize(dxf: Path, sp: Dict[str, Any], figure: str = "", number: int = 0) -> Page:
    """Render the DXF's capture window to a gray raster of exactly the placed size at ``dpi``."""
    import numpy as np
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    from ezdxf.addons.drawing import Frontend, RenderContext
    from ezdxf.addons.drawing.config import BackgroundPolicy, ColorPolicy, Configuration
    from ezdxf.addons.drawing.matplotlib import MatplotlibBackend
    from patent_figure import sheet as _sheet

    _sheet._preview_cjk_font()
    doc = read_doc(dxf)
    msp = doc.modelspace()
    win = sheet_window(msp, sp["area_mm"])
    pl = place(win, sp)
    dpi = int(sp["dpi"])
    wpx = max(1, int(round(pl["w"] / 25.4 * dpi)))
    hpx = max(1, int(round(pl["h"] / 25.4 * dpi)))
    # Agg truncates the figure size to whole pixels, and wpx/dpi*dpi can come back as wpx-1e-12.
    # Make the figure half a pixel larger and pin the axes to EXACTLY wpx × hpx pixels from the
    # bottom-left, so the window → pixel mapping is exact and the canvas is wpx × hpx.
    fig = Figure(figsize=((wpx + 0.5) / float(dpi), (hpx + 0.5) / float(dpi)), dpi=dpi)
    canvas = FigureCanvasAgg(fig)
    ax = fig.add_axes([0.0, 0.0, wpx / (wpx + 0.5), hpx / (hpx + 0.5)])
    ax.set_axis_off()
    cfg = Configuration(background_policy=BackgroundPolicy.WHITE,
                        color_policy=ColorPolicy.BLACK)
    Frontend(RenderContext(doc), MatplotlibBackend(ax), config=cfg).draw_layout(
        msp, finalize=False)
    ax.set_xlim(win[0], win[2])
    ax.set_ylim(win[1], win[3])
    ax.set_aspect("auto")
    canvas.draw()
    rgba = np.asarray(canvas.buffer_rgba())
    if rgba.shape[0] >= hpx and rgba.shape[1] >= wpx:
        rgba = rgba[rgba.shape[0] - hpx:, :wpx]          # keep the bottom-left wpx × hpx
    if rgba.shape[0] != hpx or rgba.shape[1] != wpx:
        raise DeliverableError("栅格尺寸 %s 与版式 %dx%d 不符" % (rgba.shape[:2], hpx, wpx))
    gray = ((rgba[..., 0].astype(np.uint32) * 299 + rgba[..., 1].astype(np.uint32) * 587
             + rgba[..., 2].astype(np.uint32) * 114) // 1000).astype(np.uint8)
    return Page(figure, number, dxf, gray, win, pl, dpi)


# --------------------------------------------------------------------------- writers


def write_pdf(pages: Sequence[Page], path: Path, title: str = "说明书附图") -> Path:
    """One A4 page per figure, the raster placed at its page rectangle (native resolution:
    ``interpolation='none'`` makes the PDF backend embed the image unsampled)."""
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.figure import Figure

    path = Path(path)
    tmp = path.with_name(path.name + ".tmp")
    pw, ph = PAGE_MM
    try:
        with PdfPages(str(tmp), metadata={"Title": title, "Creator": "patent-figure-kit",
                                          "Producer": "patent-figure-kit", "CreationDate": None,
                                          "ModDate": None}) as pdf:
            for page in pages:
                pl = page.placement
                fig = Figure(figsize=(pw / 25.4, ph / 25.4), dpi=72)
                ax = fig.add_axes([pl["x"] / pw, (ph - pl["top"] - pl["h"]) / ph,
                                   pl["w"] / pw, pl["h"] / ph])
                ax.set_axis_off()
                ax.imshow(page.gray, cmap="gray", vmin=0, vmax=255, interpolation="none",
                          aspect="auto")
                pdf.savefig(fig)
        tmp.replace(path)
    finally:                                  # a failed write never leaves a .tmp in the bundle
        if tmp.exists():
            tmp.unlink()
    return path


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def write_docx(pages: Sequence[Page], path: Path, title: str = "说明书附图") -> Path:
    """Word: A4, the same margins, one centred inline picture per page (page break before every
    page after the first), picture at its exact placed size. python-docx is optional."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_MM[0]), Mm(PAGE_MM[1])
    sec.top_margin, sec.left_margin = Mm(MARGINS_MM["top"]), Mm(MARGINS_MM["left"])
    sec.bottom_margin, sec.right_margin = Mm(MARGINS_MM["bottom"]), Mm(MARGINS_MM["right"])
    for i, page in enumerate(pages):
        para = doc.paragraphs[0] if (i == 0 and doc.paragraphs) else doc.add_paragraph()
        pf = para.paragraph_format
        # Every layout-relevant property is set EXPLICITLY (direct formatting beats the template's
        # docDefaults, which carry 1.15 line spacing + 10 pt after): centred, no indent, 0 before /
        # after, single line spacing — the exact structure deliverable_check accepts.
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Mm(0)
        pf.right_indent = Mm(0)
        pf.first_line_indent = Mm(0)
        if i > 0:
            pf.page_break_before = True
        run = para.add_run()
        shape = run.add_picture(io.BytesIO(page.png_bytes()),
                                width=Mm(page.placement["w"]), height=Mm(page.placement["h"]))
        try:                                     # alt text: 图 N (screen readers / search)
            shape._inline.docPr.set("descr", "图%d" % page.number)
        except Exception:
            pass
    doc.core_properties.title = title
    path = Path(path)
    tmp = path.with_name(path.name + ".tmp")
    try:
        doc.save(str(tmp))
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink()
    return path
