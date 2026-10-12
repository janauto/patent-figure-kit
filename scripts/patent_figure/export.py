"""The single export service (WP1 R1e #3/#7).

Browser ``/api/export``, REST ``POST /v1/projects/{id}/export`` and MCP ``get_export`` all export
through this module — one way to choose, verify, pin, read and release a version:

* ``acquire(ws_root, rid=None) -> Pin`` — under the version store's ``state_lock`` (the lock
  ``reclaim`` decides under) resolve the version (``None`` → current), require it to be
  **publishable** (integrity-verified against its declared artefact list, sealed ``ok``,
  result ``ok`` — a missing / truncated / zero-byte required artefact is refused) and register
  a reference, all in ONE critical section. There is no window between "validated" and
  "pinned" in which a concurrent reclaim could delete the source (R1e #7). Raises
  ``ExportRefused`` (→ HTTP 409) otherwise. There is **no fallback** to ``ws.out`` or to a
  caller-supplied result: an export reads one pinned version or nothing.
* ``Pin.release()`` — idempotent; every exit path (job create/start failure, the export body's
  ``finally``) calls it, so a failed export can never leave a version pinned forever.
* ``build_bundle(pin, dest, ...)`` — reads ONLY ``pin.vdir`` (+ the workspace flowcharts when
  asked) and writes ONLY under ``dest``: rasters are rendered into the bundle, never into the
  immutable version dir. Accounting is per figure × version × format (ok / skip / fail); a
  missing or empty source is ``fail``, never silently packaged.
* ``zip_dir(dest, archive)`` — atomic (temp + ``os.replace``).

OCP-free: DXF→PNG via ``patent_figure.sheet`` (matplotlib, lazy), DWG via a caller's ``run_cli``.
"""

from __future__ import annotations

import json
import os
import shutil
import threading
import zipfile
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from patent_figure import versions as _V

SCRIPTS = Path(__file__).resolve().parents[1]

VARIANT_DIRS = {"raw": "原始出图", "annotated": "交底版_带件号表", "filing": "递交版_仅图号"}
VARIANT_SUFFIX = {"raw": "", "annotated": "_annotated", "filing": "_filing"}
VARIANT_LABEL = {"raw": "原始", "annotated": "交底版", "filing": "递交版"}
#: REST/MCP ``variant`` parameter → the variants it packages.
REST_VARIANTS = {"filing": ("filing",), "disclosure": ("annotated",),
                 "both": ("annotated", "filing")}
DOCX_NAME = "说明书附图_交底版.docx"
NUMERALS_NAME = "reference-numerals.json"
FLOW_DIR = "流程图"


class ExportRefused(Exception):
    """No exportable version: missing, unsealed, failing verify, or not ok. → HTTP 409."""


class Pin:
    """A verified, reference-counted hold on one immutable version directory."""

    def __init__(self, store: "_V.VersionStore", rid: str, vdir: Path) -> None:
        self.store = store
        self.rid = str(rid)
        self.vdir = Path(vdir)
        self.result: Dict[str, Any] = store.read_result(self.rid) or {}
        self._released = False
        self._lock = threading.Lock()

    def release(self) -> None:
        with self._lock:
            if self._released:
                return
            self._released = True
        self.store.release(self.rid)

    @property
    def released(self) -> bool:
        return self._released

    def __enter__(self) -> "Pin":
        return self

    def __exit__(self, *exc) -> None:
        self.release()


def acquire(ws_root, rid: Optional[str] = None) -> Pin:
    """Verify + pin a version atomically (see module docstring). ``rid=None`` → current."""
    vs = _V.VersionStore(ws_root)
    try:
        rid, vdir = vs.acquire_verified_version(rid)
    except ValueError as exc:
        raise ExportRefused(str(exc))
    try:
        return Pin(vs, rid, vdir)
    except BaseException:
        vs.release(rid)
        raise


def _nonempty(path: Path) -> bool:
    try:
        return path.is_file() and path.stat().st_size > 0
    except OSError:
        return False


def _png(src_dxf: Path, png: Path, dpi: int) -> None:
    from patent_figure import sheet as _sheet       # matplotlib is imported lazily inside;
    _sheet.render_preview(src_dxf, png, dpi=dpi)    # it holds the process-wide lock (R2m)


def write_docx(pngs: Sequence[Path], path: Path) -> None:
    """说明书附图.docx：每图一页，A4 版心等比缩放（python-docx + Pillow，可选依赖）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    from PIL import Image
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.left_margin = Cm(2.5), Cm(2.5)
    sec.bottom_margin, sec.right_margin = Cm(1.5), Cm(1.5)
    max_w, max_h = 17.0, 25.0
    for i, png in enumerate(pngs):
        with Image.open(png) as im:
            w, h = im.size
        width = min(max_w, max_h * w / h)
        para = doc.paragraphs[0] if i == 0 and doc.paragraphs else doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(png), width=Cm(width))
        if i < len(pngs) - 1:
            doc.add_page_break()
    doc.save(str(path))


def _dwg(src_dxf: Path, sub: Path, run_cli: Callable, keep_dxf: bool) -> Tuple[bool, str]:
    """Convert one DXF to DWG inside the bundle folder ``sub`` (never next to the version's
    file). AutoCAD Core first, LibreDWG as fallback."""
    local = sub / src_dxf.name
    copied = False
    if not local.is_file():
        shutil.copy2(src_dxf, local)
        copied = True
    dwg = local.with_suffix(".dwg")
    try:
        proc = run_cli([str(SCRIPTS / "autocad_core_dxf_to_dwg.py"), str(local), str(dwg)],
                       timeout=600)
        how = "AutoCAD 转换成功"
        if proc.returncode != 0 or not dwg.is_file():
            proc = run_cli([str(SCRIPTS / "libredwg_dxf_to_dwg.py"), str(local), "-o", str(sub)],
                           timeout=600)
            how = "AutoCAD 失败，改用 LibreDWG（exit=%d）" % proc.returncode
    finally:
        if copied and not keep_dxf:
            try:
                local.unlink()
            except OSError:
                pass
    return _nonempty(dwg), how


DOC_BASENAME = "说明书附图_%s.%s"           # % (递交版|交底版, pdf|docx)
FAILED_DIR = "_未通过核验"
REPORT_NAME = "核验报告.json"
CHECKLIST_NAME = "核对清单.txt"
DESCRIPTION_NAME = "附图说明.txt"
#: Variants that carry the figure number 「图N」 and therefore can be a filing document.
DOC_VARIANTS = ("annotated", "filing")


def _fig_rows(figs: Sequence[dict], vdir: Path, suf: str) -> List[Dict[str, Any]]:
    rows = []
    for i, f in enumerate(figs, start=1):
        rows.append({"figure": str(f["id"]), "number": int(f.get("number") or i),
                     "dxf": vdir / (str(f["id"]) + suf + ".dxf")})
    return rows


def _build_document(variant: str, fmt: str, rows: List[Dict[str, Any]], sp: Dict[str, Any],
                    dest: Path, pages_cache: Dict[str, list]) -> Dict[str, Any]:
    """Write ONE final document (PDF or Word) for a variant, then check it page by page from the
    written file (``deliverable_check``). EVERY non-ok outcome — a failed check, a checker that
    raised, a write that raised after leaving a file behind — ends in the same place: the file
    (if any) is moved to ``_未通过核验/`` and ``path`` records where it really is. Only a passing
    document stays at the package root (WP2 R2 #4). Returns the check report + ``path``."""
    from patent_figure import deliverable as _D
    from patent_figure import deliverable_check as _C
    name = DOC_BASENAME % (VARIANT_LABEL[variant], fmt)
    title = "说明书附图（%s）" % VARIANT_LABEL[variant]
    out = dest / name
    layout = {"layout_verified": "structural" if fmt == "docx" else "file-geometry"}

    def failed(errors: List[str]) -> Dict[str, Any]:
        return dict({"file": name, "format": fmt, "ok": False, "status": "fail", "pages": 0,
                     "expected_pages": len(rows), "page_checks": [], "errors": errors}, **layout)

    missing = [r for r in rows if not _nonempty(Path(r["dxf"]))]
    if missing:
        report = failed(["版本内源图缺失或为空：%s" % "、".join(
            Path(r["dxf"]).name for r in missing)])
    else:
        report = None
        try:
            if variant not in pages_cache:
                pages_cache[variant] = [_D.rasterize(Path(r["dxf"]), sp, r["figure"],
                                                     r["number"]) for r in rows]
            if fmt == "pdf":
                _D.write_pdf(pages_cache[variant], out, title)
            else:
                _D.write_docx(pages_cache[variant], out, title)
        except Exception as exc:
            report = failed(["生成失败：%s" % exc])
        if report is None:
            # R2j ruling — HASH SEAL: the digest is taken the moment the file is written and the
            # check runs against it, so the file checked is the file written (and the digest
            # recorded in the export result is the one every later re-check must match)
            sha = _C.sha256_of(out)
            try:
                report = _C.check_document(out, _C.sources_from(rows, sp["area_mm"]), sp,
                                           expected_sha256=sha)
            except Exception as exc:                # an unverifiable document is a failed one
                report = failed(["核验程序异常，无法确认这份文件合格：%s: %s"
                                 % (type(exc).__name__, exc)])
            report["sha256"] = sha
    # only a PASSING document stays at the root: ``fail`` and ``unverified`` (no renderer — R2j
    # ruling) both go to _未通过核验/, with the reason
    report["path"] = name if report.get("ok") else None
    if not report.get("ok") and out.exists():
        (dest / FAILED_DIR).mkdir(parents=True, exist_ok=True)
        os.replace(str(out), str(dest / FAILED_DIR / name))
        report["path"] = "%s/%s" % (FAILED_DIR, name)
    return report


def _doc_reason(rep: Dict[str, Any], fid: str) -> str:
    page = next((p for p in rep.get("page_checks", []) if p.get("figure") == fid), None)
    errs = list(rep.get("errors") or []) + list((page or {}).get("errors") or [])
    if not errs:          # this page is fine, but the document as a whole failed
        others = ["第 %d 页：%s" % (p.get("page"), "；".join(p.get("errors") or []))
                  for p in rep.get("page_checks", []) if not p.get("ok")]
        errs = ["本页通过，但同一文件未通过核验，整份文件不能递交（%s）" % "；".join(others[:2])]
    return "；".join(errs[:4])


def checklist_text(result: Dict[str, Any], bundle: Dict[str, Any], sp: Optional[dict]) -> str:
    """《核对清单.txt》：机器核验结果 + 必须由人核对的事项（M1 无说明书自动对表）。"""
    lines = ["说明书附图导出包 · 核对清单", "",
             "渲染版本：%s" % bundle.get("render_job_id", ""), ""]
    if sp:
        lines += ["版式：A4，页边 上%g/左%g/下%g/右%g mm，图区 %s（%g×%g mm），%d dpi；"
                  "缩放后字高不低于 %g mm、栅格不低于 %g px。" % (
                      sp["margins_mm"]["top"], sp["margins_mm"]["left"],
                      sp["margins_mm"]["bottom"], sp["margins_mm"]["right"], sp["preset"],
                      sp["area_mm"][0], sp["area_mm"][1], sp["dpi"], sp["min_text_mm"],
                      sp["min_text_px"]), ""]
    lines.append("一、机器逐页核验（文件本身，不是源图）")
    docs = bundle.get("documents") or {}
    if not docs:
        lines.append("  （本次未生成 PDF / Word）")
    if any(r.get("layout_verified") == "structural" for r in docs.values()):
        lines.append("  注：Word 按文档结构核验，实际排版以 Word 打开为准（PDF 按文件里的页面几何核验）。")
    for name, rep in sorted(docs.items()):
        unverified = rep.get("status") == "unverified"
        lines.append("  %s %s：%d/%d 页%s" % (
            "[通过]" if rep.get("ok") else ("[未核验]" if unverified else "[不通过]"), name,
            max(rep.get("pages", 0), 0), rep.get("expected_pages", 0),
            "" if rep.get("ok") else "——已移到 %s/，%s。原因：%s" % (
                FAILED_DIR, "未经可见性核验，不要递交" if unverified else "不要递交",
                "；".join((rep.get("errors") or [])[:2]
                         + [e for p in rep.get("page_checks", [])
                            for e in p.get("errors", [])][:2]))))
    skipped = ["%s：%s" % (k, v) for k, v in sorted((bundle.get("reasons") or {}).items())
               if (bundle.get("manifest") or {}).get(k) == "skip"]
    if skipped:
        lines.append("  跳过：" + "；".join(skipped[:6]))
    lines += ["", "二、必须人工核对（本工具 M1 不做说明书自动对表）",
              "  [ ] 说明书「附图标记说明」与每张图上的标号逐条一致（见 reference-numerals.json）",
              "  [ ] 说明书「附图说明」段与本包《附图说明.txt》一致，图号连续（图1、图2…）",
              "  [ ] 每张图上没有内部件号、尺寸、明细表等不该出现的文字",
              "  [ ] 递交用的是「递交版」（仅图号）；「交底版」带件号表，只用于内部交底",
              "  [ ] DXF/DWG 是源文件，供再编辑；递交请用 PDF 或 Word", "",
              "三、提示", "  %s。" % (sp or {}).get("note", "精确图片参数以当期《电子申请文件格式要求说明》为准"),
              "  核验只对导出包原文件有效：每个文件的 sha256 记在《核验报告.json》与导出记录里，"
              "文件被改动（哪怕一个字节）即判「文件已被修改」，须用导出包原文件递交。",
              "  本工具按 CNIPA 形式要求排版，最终以受理为准，不构成法律意见。", ""]
    return "\n".join(lines)


def build_bundle(pin: Pin, dest: Path, *, variants: Iterable[str] = ("annotated", "filing"),
                 formats: Iterable[str] = ("docx", "dxf"),
                 raw_formats: Optional[Iterable[str]] = None, dpi: int = 300,
                 run_cli: Optional[Callable] = None,
                 flows: Iterable[Tuple[str, Path, dict]] = (),
                 description_txt: bool = False, sheet_preset: str = "standard",
                 checklist_txt: bool = False,
                 dwg_available: Optional[bool] = None) -> Dict[str, Any]:
    """Build an export bundle from ONE pinned version into ``dest``.

    ``variants`` ⊆ {raw, annotated, filing}; ``formats`` ⊆ {dxf, svg, png, docx, dwg, pdf}
    (``raw_formats`` overrides the list for the raw variant). ``flows`` = ``(id, out_dir,
    result)`` of flowcharts copied under 流程图/ (workspace-level, not version-bound — they are
    NOT pages of the PDF/Word deliverables, so 页数 = 结构附图数).

    F1 (WP2): ``pdf`` / ``docx`` are the PRIMARY deliverables — one document per numbered variant
    (``说明书附图_递交版.pdf`` …), A4 pages at ``dpi`` under ``sheet_preset``, each checked page by
    page from the written file (``deliverable_check``); a failing document goes to
    ``_未通过核验/`` and its items are ``fail`` with the reason. A PDF passes only after an
    independent render (R2j ruling); with no renderer it is ``unverified`` — also moved to
    ``_未通过核验/``, never ok. DWG is ``skip`` (with reason) when no converter is available
    (``dwg_available``; probed when None).

    Returns ``{manifest, reasons, items, documents, files, notes, dwg_log, figure_descriptions,
    spec, sha256, ok}`` where ``manifest`` maps ``<fig>/<版本>/<fmt>`` → ``ok`` / ``skip`` /
    ``fail`` / ``unverified`` (``reasons`` explains every non-ok item), ``sha256`` maps every
    bundle file except 《核验报告.json》 to its digest as written (the HASH SEAL: a later re-check
    of a bundle file compares against it first) and ``ok`` is True iff no requested item failed
    or is unverified."""
    vdir, result = pin.vdir, pin.result
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    formats = list(formats)
    raw_formats = list(raw_formats) if raw_formats is not None else formats
    # workbench results mark a delivered figure ``pass``; direct-CLI results ``gate_pass``
    figs = sorted([f for f in result.get("figures", [])
                   if (f.get("pass") if "pass" in f else f.get("gate_pass"))],
                  key=lambda f: (f.get("number") or 999, str(f.get("id"))))
    manifest: Dict[str, str] = {}
    reasons: Dict[str, str] = {}
    notes: List[str] = []
    dwg_log: List[str] = []
    sp = None
    if any(f in ("pdf", "docx") for f in formats):
        from patent_figure import deliverable as _D
        try:
            sp = _D.spec(sheet_preset, dpi)
        except _D.DeliverableError as exc:
            notes.append(str(exc))
    if "dwg" in formats + raw_formats and dwg_available is None:
        from patent_figure import capabilities as _CAP
        dwg_available = _CAP.dwg_route() is not None

    def put(key: str, status: str, reason: str = "") -> None:
        manifest[key] = status
        if reason:
            reasons[key] = reason
        elif key in reasons:
            del reasons[key]

    for fig in figs:
        fid = str(fig["id"])
        for variant in variants:
            suf, label = VARIANT_SUFFIX[variant], VARIANT_LABEL[variant]
            sub = dest / VARIANT_DIRS[variant]
            src_dxf = vdir / (fid + suf + ".dxf")
            fmts = raw_formats if variant == "raw" else formats
            for fmt in fmts:
                key = "%s/%s/%s" % (fid, label, fmt)
                if fmt in ("pdf", "docx"):
                    if variant not in DOC_VARIANTS:
                        put(key, "skip", "原始出图没有图号，不出 PDF/Word")
                    continue                                 # documents: second pass below
                if fmt == "dwg" and variant == "raw":
                    put(key, "skip", "原始出图不转 DWG")
                    continue
                src = vdir / (fid + suf + (".svg" if fmt == "svg" else ".dxf"))
                if not _nonempty(src):                           # never package a hole
                    put(key, "fail", "版本内源文件缺失或为空（%s）" % src.name)
                    notes.append("%s：版本内源文件缺失或为空（%s）" % (key, src.name))
                    continue
                sub.mkdir(parents=True, exist_ok=True)
                if fmt in ("dxf", "svg"):
                    shutil.copy2(src, sub / src.name)
                    if _nonempty(sub / src.name):
                        put(key, "ok")
                    else:
                        put(key, "fail", "复制后文件为空")
                elif fmt == "png":
                    png = sub / (fid + suf + ".png")
                    try:
                        if not _nonempty(png):
                            _png(src_dxf, png, dpi)
                    except Exception as exc:
                        notes.append("%s：PNG 生成失败：%s" % (key, exc))
                    if _nonempty(png):
                        put(key, "ok")
                    else:
                        put(key, "fail", "PNG 生成失败")
                elif fmt == "dwg":
                    if run_cli is None or not dwg_available:
                        put(key, "skip", "未检测到 DWG 转换器（AutoCAD Core Console 或 LibreDWG）"
                            if not dwg_available else "本入口未接 DWG 转换")
                        continue
                    ok, how = _dwg(src_dxf, sub, run_cli, keep_dxf="dxf" in fmts)
                    put(key, "ok" if ok else "fail", "" if ok else "DWG 转换失败：%s" % how)
                    dwg_log.append("%s/%s: %s" % (VARIANT_DIRS[variant], src_dxf.stem,
                                                  how if ok else how + "，转换失败"))
                else:
                    put(key, "skip", "未知格式")
                    notes.append("%s：未知格式" % key)

    # ---- F1 primary deliverables: one PDF / Word per numbered variant, checked page by page ---
    documents: Dict[str, Dict[str, Any]] = {}
    pages_cache: Dict[str, list] = {}
    for variant in [v for v in variants if v in DOC_VARIANTS]:
        label = VARIANT_LABEL[variant]
        for fmt in [f for f in ("pdf", "docx") if f in formats]:
            keys = ["%s/%s/%s" % (str(f["id"]), label, fmt) for f in figs]
            if not figs:
                continue
            if sp is None:
                for k in keys:
                    put(k, "fail", "版式参数无效：%s" % "；".join(notes[-1:]))
                continue
            if fmt == "docx":
                from patent_figure import deliverable as _D
                if not _D.docx_available():
                    for k in keys:
                        put(k, "skip", "未安装 python-docx，不生成 Word（pip install python-docx）")
                    continue
            rep = _build_document(variant, fmt, _fig_rows(figs, vdir, VARIANT_SUFFIX[variant]),
                                  sp, dest, pages_cache)
            documents[rep["file"]] = rep
            for f, k in zip(figs, keys):
                page = next((p for p in rep.get("page_checks", [])
                             if p.get("figure") == str(f["id"])), None)
                good = rep["ok"] and page is not None and page.get("ok")
                status = "ok" if good else (
                    "unverified" if rep.get("status") == "unverified" else "fail")
                put(k, status, "" if good else _doc_reason(rep, str(f["id"])))
    pages_cache.clear()

    flows = [(fid, Path(odir), res or {}) for fid, odir, res in flows]
    for fid, odir, _res in flows:
        for ext in (".dxf", ".png", ".svg"):
            src = odir / (fid + ext)
            if _nonempty(src):
                (dest / FLOW_DIR).mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, dest / FLOW_DIR / src.name)

    numerals_src = vdir / NUMERALS_NAME
    numerals_doc: Dict[str, Any] = {}
    if _nonempty(numerals_src):
        shutil.copy2(numerals_src, dest / NUMERALS_NAME)
        try:
            numerals_doc = json.loads(numerals_src.read_text(encoding="utf-8"))
        except ValueError:
            numerals_doc = {}
    desc = [(f.get("annotation") or {}).get("figure_description", "") for f in figs]
    desc += [(res or {}).get("figure_description", "") for _fid, _odir, res in flows]
    desc = [d for d in desc if d]
    if desc:
        desc[-1] = desc[-1].rstrip("；") + "。"
    if description_txt:
        text = ("【附图说明】\n" + "\n".join(desc) + "\n\n"
                + (numerals_doc.get("description_zh") or "") + "\n")
        (dest / DESCRIPTION_NAME).write_text(text, encoding="utf-8")

    bundle = {"manifest": manifest, "reasons": reasons, "documents": documents,
              "render_job_id": pin.rid}
    if checklist_txt:
        (dest / CHECKLIST_NAME).write_text(checklist_text(result, bundle, sp), encoding="utf-8")

    # ---- 文件齐全：every item reported ok must be a non-empty file in the bundle -------------
    from patent_figure import deliverable_check as _C
    expected_files = []
    for k, v in manifest.items():
        if v != "ok":
            continue
        fid, label, fmt = k.split("/")
        variant = next(x for x, y in VARIANT_LABEL.items() if y == label)
        if fmt in ("pdf", "docx"):
            expected_files.append(DOC_BASENAME % (label, fmt))
        elif fmt in ("dxf", "svg", "png", "dwg"):
            expected_files.append("%s/%s%s.%s" % (VARIANT_DIRS[variant], fid,
                                                  VARIANT_SUFFIX[variant], fmt))
    if description_txt:
        expected_files.append(DESCRIPTION_NAME)
    if checklist_txt:
        expected_files.append(CHECKLIST_NAME)
    missing_files = _C.check_files(dest, sorted(set(expected_files)))
    for k, v in list(manifest.items()):
        fid, label, fmt = k.split("/")
        variant = next(x for x, y in VARIANT_LABEL.items() if y == label)
        f = DOC_BASENAME % (label, fmt) if fmt in ("pdf", "docx") else \
            "%s/%s%s.%s" % (VARIANT_DIRS[variant], fid, VARIANT_SUFFIX[variant], fmt)
        if v == "ok" and f in missing_files:
            put(k, "fail", "包内缺文件或为空：%s" % f)

    # ---- HASH SEAL (R2j ruling): sha256 of every delivered file, as written ---------------
    # (the report cannot hold its own digest; every other file of the bundle is in the map)
    sha256 = seal_files(dest)
    for name, rep in documents.items():
        at = rep.get("path") or ("%s/%s" % (FAILED_DIR, name))
        if rep.get("sha256") and sha256.get(at) != rep["sha256"]:
            for k in [k for k in manifest if k.endswith("/" + rep["format"])
                      and DOC_BASENAME % (k.split("/")[1], rep["format"]) == name]:
                put(k, "fail", "%s 在打包过程中被改动（sha256 与写出时不符）" % name)
    report = {"schema": "patent-figure-export-check/1", "render_job_id": pin.rid,
              "authority": AUTHORITY_NOTE,
              "spec": sp, "documents": documents, "missing_files": missing_files,
              "manifest": manifest, "reasons": reasons, "sha256": sha256,
              "sha256_note": SHA256_NOTE}
    (dest / REPORT_NAME).write_text(json.dumps(report, ensure_ascii=False, indent=2,
                                               default=str), encoding="utf-8")
    files = sorted(p.relative_to(dest).as_posix() for p in dest.rglob("*") if p.is_file())
    return {"manifest": manifest, "reasons": reasons,
            "items": items_of(manifest, reasons, documents),
            "documents": documents, "missing_files": missing_files, "files": files,
            "notes": notes, "dwg_log": dwg_log, "figure_descriptions": desc,
            "render_job_id": pin.rid, "spec": sp, "sha256": sha256,
            "ok": not any(v in ("fail", "unverified") for v in manifest.values())}


SHA256_NOTE = ("核验只对导出包原文件有效：sha256 在写出每个文件时记录（本报告自身除外）；"
               "复核时先比对，不一致即判「文件已被修改，请使用导出包原文件」，不再做其余核验。"
               "权威记录是服务端导出结果（result.sha256），包内本报告仅供对照。")


AUTHORITY_NOTE = ("以服务端封口记录为准：本报告随包分发、可被一并改动，只作对照；权威记录是服务端导出目录里的"
                  " seal.json（封住本包全部文件、bundle.zip、导出报告与作业结果），其 sha256 记在导出作业"
                  "结果 result.seal.sha256")
SEAL_NAME = "seal.json"
RESULT_NAME = "result.json"


def seal_export(export_dir: Path, result: Dict[str, Any]) -> Dict[str, str]:
    """R2k — seal an export's OWN directory (the service's ``<jobs>/<export_id>/``): write the job
    ``result`` (manifest, reasons, documents, sha256 …) as ``result.json``, then ``seal.json`` =
    the export report (ok / manifest / reasons / documents / bundle sha256) + ``{path: {size,
    sha256}}`` of EVERY file in the directory (``bundle/**``, ``bundle.zip``, ``result.json``),
    fsynced; then every file is made read-only. Returns ``{file, sha256}`` — the digest of
    ``seal.json`` goes into the job result, so the persisted job record and the sealed directory
    vouch for each other. The bundle's own 《核验报告.json》 is only a copy for reference."""
    import hashlib
    from patent_figure import deliverable_check as _C
    from patent_figure import project_io as _pio
    export_dir = Path(export_dir)
    _pio.atomic_write_json(export_dir / RESULT_NAME, result)
    files: Dict[str, Dict[str, Any]] = {}
    for p in sorted(export_dir.rglob("*")):
        rel = p.relative_to(export_dir).as_posix()
        if p.is_file() and rel != SEAL_NAME:
            files[rel] = {"size": p.stat().st_size, "sha256": _C.sha256_of(p)}
    seal = {"schema": "patent-figure-export-seal/1",
            "export_id": result.get("export_id"), "render_job_id": result.get("render_job_id"),
            "ok": result.get("ok"), "manifest": result.get("manifest"),
            "reasons": result.get("reasons"), "documents": result.get("documents"),
            "sha256": result.get("sha256"), "files": files}
    data = json.dumps(seal, ensure_ascii=False, indent=1, sort_keys=True, default=str)
    target = export_dir / SEAL_NAME
    tmp = target.with_name("%s.%d.%s.tmp" % (target.name, os.getpid(), os.urandom(4).hex()))
    try:
        with open(str(tmp), "w", encoding="utf-8") as fh:
            fh.write(data)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(str(tmp), str(target))
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    for p in export_dir.rglob("*"):
        if p.is_file():
            try:
                os.chmod(str(p), 0o444)
            except OSError:
                pass
    return {"file": SEAL_NAME,
            "sha256": hashlib.sha256((export_dir / SEAL_NAME).read_bytes()).hexdigest()}


def verify_export_seal(export_dir: Path, seal_sha256: Optional[str]) -> List[str]:
    """Errors (empty = intact) of a sealed export directory against the digest of its seal kept in
    the job result: the seal itself unchanged, every sealed file present with its size and
    sha256, and no file that the seal does not list."""
    import hashlib
    from patent_figure import deliverable_check as _C
    export_dir = Path(export_dir)
    try:
        raw = (export_dir / SEAL_NAME).read_bytes()
    except OSError:
        return ["导出目录缺封口记录 %s" % SEAL_NAME]
    if seal_sha256 and hashlib.sha256(raw).hexdigest() != seal_sha256:
        return ["封口记录 %s 与导出作业里记录的 sha256 不符（被改动）" % SEAL_NAME]
    try:
        files = json.loads(raw.decode("utf-8"))["files"]
    except (ValueError, KeyError, UnicodeDecodeError):
        return ["封口记录 %s 读不了" % SEAL_NAME]
    errors = []
    for rel, info in sorted(files.items()):
        p = export_dir / rel
        if not p.is_file():
            errors.append("封口文件缺失：%s" % rel)
        elif p.stat().st_size != info.get("size") or _C.sha256_of(p) != info.get("sha256"):
            errors.append("文件已被修改：%s（与封口记录的 sha256 不符）" % rel)
    present = {p.relative_to(export_dir).as_posix() for p in export_dir.rglob("*") if p.is_file()}
    extra = sorted(present - set(files) - {SEAL_NAME})
    if extra:
        errors.append("导出目录里有封口后多出的文件：%s" % "、".join(extra[:5]))
    return errors


def seal_files(dest: Path) -> Dict[str, str]:
    """``{bundle-relative path: sha256}`` of every file under ``dest`` except 《核验报告.json》."""
    from patent_figure import deliverable_check as _C
    out: Dict[str, str] = {}
    for p in sorted(Path(dest).rglob("*")):
        rel = p.relative_to(dest).as_posix()
        if p.is_file() and rel != REPORT_NAME:
            out[rel] = _C.sha256_of(p) or ""
    return out


def items_of(manifest: Dict[str, str], reasons: Dict[str, str],
             documents: Optional[Dict[str, Dict[str, Any]]] = None) -> List[Dict[str, Any]]:
    """Per (版本, 格式) summary rows for the export screen: status = fail if any figure failed,
    ok if every figure is ok, otherwise skip; ``reason`` = the first reason found."""
    groups: Dict[Tuple[str, str], List[str]] = {}
    for k in manifest:
        _fid, label, fmt = k.split("/")
        groups.setdefault((label, fmt), []).append(k)
    out = []
    for (label, fmt), keys in sorted(groups.items()):
        sts = [manifest[k] for k in keys]
        status = ("fail" if "fail" in sts else "unverified" if "unverified" in sts else
                  "ok" if all(s == "ok" for s in sts) else "skip")
        reason = next((reasons[k] for k in keys if manifest[k] == status and k in reasons), "")
        doc = DOC_BASENAME % (label, fmt) if fmt in ("pdf", "docx") else None
        if doc is not None:                      # where the file really is (failed → 未通过核验/)
            doc = ((documents or {}).get(doc) or {}).get("path") or (
                doc if status == "ok" else None)
        out.append({"variant": label, "format": fmt, "status": status, "reason": reason,
                    "figures": len(keys), "file": doc, "role": (
                        "主交付" if fmt in ("pdf", "docx") else
                        ("源文件，供再编辑" if fmt in ("dxf", "dwg") else "预览"))})
    return out


def zip_dir(src: Path, archive: Path) -> Path:
    """Zip ``src`` into ``archive`` atomically (temp + os.replace). No ``shutil.make_archive``:
    it calls ``os.getcwd()``, which raises when the service's cwd is unreadable."""
    src, archive = Path(src), Path(archive)
    tmp = archive.with_name("%s.%d.%s.tmp" % (archive.name, os.getpid(), os.urandom(4).hex()))
    try:
        with zipfile.ZipFile(str(tmp), "w", zipfile.ZIP_DEFLATED) as zf:
            for f in sorted(src.rglob("*")):
                if f.is_file():
                    zf.write(str(f), f.relative_to(src).as_posix())
        os.replace(str(tmp), str(archive))
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    return archive
