"""Plan Studio 扩展：规范性标注、流程图、AI 起草、说明书附图 Word 导出。

都挂在 ``plan_studio.build_app`` 上，单工程（``plan_studio.py ASM.stp``）与多工程
工作台（``workbench.py``）共用。原则不变：出图与编号全部走 skill 的 CLI；大模型只写
语义（零件中文名、流程图的步骤与连线），写回 plan 前都过校验，人填过的内容绝不覆盖。

图号规则：结构附图按 plan.figures 的顺序为 图1…图n，流程图接着编 图n+1…。
"""

from __future__ import annotations

import json
import re
import shutil
import threading
import time
from pathlib import Path
from types import SimpleNamespace
from typing import Callable, Dict, List, Optional

RENDER_CLI = Path(__file__).resolve().parent / "render_patent_figure.py"

SCRIPTS = Path(__file__).resolve().parent
ANNOTATE_CLI = SCRIPTS / "annotate_figure_sheet.py"
FLOW_CLI = SCRIPTS / "render_flowchart.py"
FLOW_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,31}$")
VARIANTS = ("annotated", "filing")
VARIANT_LABEL = {"annotated": "交底版（图号 + 件号名称表）", "filing": "递交版（仅图号）",
                 "raw": "原始出图（描述性图题）"}

#: PNG preview resolution at export (WP1-5): the interactive path previews via SVG only, so
#: raster PNGs are produced just once, here, at export time. 300 dpi matches the F1 default.
EXPORT_PNG_DPI = 300


def _load(path: Path, default=None):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def _dump(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


# --------------------------------------------------------------------------- #
# 结构附图：出图后自动出交底版 / 递交版                                          #
# --------------------------------------------------------------------------- #

def render_version(vd: Path, plan_path: Path, assembly_path: Path, cache: Path,
                   run_cli: Callable, svg_of: Callable) -> tuple:
    """Render + annotate **into an immutable version directory** ``vd`` and build the result +
    manifest, WITHOUT publishing (the caller commits under the state lock). Shared by the
    workbench job and the plan_studio sub-app so there is exactly one render path (R1b #5).

    Success is decided from the renderer's own ``--json`` report, not a directory scan: a figure
    counts as delivered only if the report marks it ``gate_pass``; ``ok`` requires the process to
    exit 0, the report to pass, and **every** figure the frozen plan expects to be delivered
    (R1b #1 — a partial failure can no longer read as success)."""
    report_path = vd / "render-report.json"
    proc = run_cli([str(RENDER_CLI), str(plan_path), "--assembly", str(assembly_path),
                    "-o", str(vd), "--cache", str(cache), "--json", str(report_path)])
    log = ((proc.stdout or "")[-6000:] + "\n" + (proc.stderr or "")[-2000:]).strip()
    report = _load(report_path, {}) or {}
    plan = _load(plan_path, {}) or {}
    expected = [str(f.get("id")) for f in plan.get("figures", []) if isinstance(f, dict)]
    delivered = [str(f.get("id")) for f in report.get("figures", []) if f.get("gate_pass")]
    ok = (proc.returncode == 0 and bool(report.get("pass"))
          and bool(expected) and set(expected) == set(delivered))

    result: dict = {"figures": [], "ok": ok, "numerals": None, "log": log,
                    "render_report": {"pass": bool(report.get("pass")),
                                      "returncode": proc.returncode,
                                      "errors": report.get("errors", [])}}
    numerals_file = vd / "reference-numerals.json"
    if numerals_file.is_file():
        result["numerals"] = _load(numerals_file, {})
    for dxf in sorted(vd.glob("*.dxf")):            # only delivered figs land in vd root
        fid = dxf.stem
        if fid.endswith(("_annotated", "_filing", "_engineering")):
            continue
        (vd / (fid + ".svg")).write_text(svg_of(dxf), encoding="utf-8")
        qa_file = vd / (fid + ".qa.json")
        qa = _load(qa_file) if qa_file.is_file() else None
        result["figures"].append({"id": fid, "pass": bool(qa and qa.get("pass")),
                                  "qa": qa, "svg": "/api/preview/%s.svg" % fid})
    present = {f["id"] for f in result["figures"]}
    for fid in expected:                            # make expected-but-missing visible as failed
        if fid not in present:
            result["figures"].append({"id": fid, "pass": False, "qa": None, "svg": None,
                                      "missing": True})
    if any(f["pass"] for f in result["figures"]):
        annotate_all(SimpleNamespace(plan=plan_path, out=vd, root=vd), run_cli, svg_of, result)
    manifest = {"ok": ok, "expected": expected, "delivered": delivered,
                "figures": [{"id": f["id"], "pass": bool(f["pass"]),
                             "variants": sorted((f.get("variants") or {}).keys())}
                            for f in result["figures"]]}
    return ok, result, manifest


def ensure_preview_png(png_path: Path, dpi: int = EXPORT_PNG_DPI) -> bool:
    """Lazily render a missing PNG from its sibling DXF (WP1-5).

    The interactive path makes SVGs only; a PNG is produced on first request here (and at
    export). Returns True if the PNG exists after the call. matplotlib is imported lazily
    inside render_preview, so a page that never asks for a PNG never pays for matplotlib.
    """
    png_path = Path(png_path)
    if png_path.is_file():
        return True
    dxf = png_path.with_suffix(".dxf")
    if not dxf.is_file():
        return False
    try:
        from patent_figure import sheet as _sheet
        _sheet.render_preview(dxf, png_path, dpi=dpi)
        return png_path.is_file()
    except Exception:
        return False


def annotate_all(ws, run_cli: Callable, svg_of: Callable, result: dict) -> None:
    """给 result["figures"] 里每张已渲染的图补上两个规范版本与附图说明。"""
    plan = _load(ws.plan, {}) or {}
    order = [f.get("id") for f in plan.get("figures", [])]
    numerals = ws.out / "reference-numerals.json"
    if not numerals.is_file():
        return
    descriptions = []
    for fig in result.get("figures", []):
        fid = fig["id"]
        if fid not in order:
            continue
        fig["number"] = order.index(fid) + 1
        if not fig.get("pass"):          # 未过 QA 的图不出规范版本（闸门不可绕）
            continue
        number = order.index(fid) + 1
        fig["number"] = number
        fig["variants"] = {"raw": "/api/preview/%s.svg" % fid}
        src = ws.out / (fid + ".dxf")
        for variant in VARIANTS:
            dst = ws.out / ("%s_%s.dxf" % (fid, variant))
            rep_path = ws.out / ("%s_%s.json" % (fid, variant))
            # WP1-5: no --preview here. The interactive preview of the annotated / filing
            # variants goes through the same hand-written interactive SVG (svg_of) as the raw
            # figure; the raster PNG is produced only at export (export_bundle), by dpi. This
            # removes one matplotlib launch per variant per render from the "出图" click.
            args = [str(ANNOTATE_CLI), str(src), "--numerals", str(numerals),
                    "--figure-number", str(number), "-o", str(dst),
                    "--json", str(rep_path)]
            if variant == "filing":
                args.append("--no-table")
            proc = run_cli(args, timeout=300)
            if proc.returncode == 0 and dst.is_file():
                (ws.out / (dst.stem + ".svg")).write_text(svg_of(dst), encoding="utf-8")
                fig["variants"][variant] = "/api/preview/%s.svg" % dst.stem
                rep = _load(rep_path, {}) or {}
                if variant == "annotated":
                    fig["annotation"] = {
                        "table": rep.get("table"), "warnings": rep.get("warnings", []),
                        "figure_description": rep.get("figure_description", ""),
                        "svg": "/api/preview/%s.svg" % dst.stem,
                        # PNG is now rendered lazily on first request (served from the sibling
                        # DXF), so this URL is valid without an eager matplotlib launch.
                        "png": "/api/preview/%s.png" % dst.stem,
                    }
                    if rep.get("figure_description"):
                        descriptions.append(rep["figure_description"])
            else:
                fig.setdefault("annotation_errors", []).append(
                    "%s: %s" % (variant, (proc.stderr or proc.stdout)[-300:]))
    result["figure_descriptions"] = descriptions
    result["variant_labels"] = VARIANT_LABEL


# --------------------------------------------------------------------------- #
# 流程图                                                                        #
# --------------------------------------------------------------------------- #

def flow_dir(ws) -> Path:
    return ws.root / "flowcharts"


def flow_out(ws) -> Path:
    return ws.root / "flow_out"


def list_flowcharts(ws) -> List[dict]:
    plan = _load(ws.plan, {}) or {}
    base = len(plan.get("figures", []))
    rows = []
    for i, path in enumerate(sorted(flow_dir(ws).glob("*.json"))):
        spec = _load(path, {}) or {}
        fid = path.stem
        res = _load(flow_out(ws) / (fid + ".result.json"), None)
        rows.append({
            "id": fid, "number": base + i + 1, "title": spec.get("title", ""),
            "spec": spec, "result": res,
            "svg": ("/api/flow-preview/%s.svg" % fid)
            if (flow_out(ws) / (fid + ".svg")).is_file() else None,
        })
    return rows


def save_flowchart(ws, fid: str, spec: dict) -> dict:
    from patent_figure import flowchart as FC
    if not FLOW_ID.match(fid):
        raise ValueError("流程图 id 只能用字母数字下划线连字符")
    _dump(flow_dir(ws) / (fid + ".json"), spec)
    return {"issues": FC.validate(spec)}


def render_flowchart(ws, fid: str, run_cli: Callable, svg_of: Callable) -> dict:
    spec_path = flow_dir(ws) / (fid + ".json")
    if not spec_path.is_file():
        raise FileNotFoundError(fid)
    number = next((r["number"] for r in list_flowcharts(ws) if r["id"] == fid), None)
    out = flow_out(ws) / (fid + ".dxf")
    res_path = flow_out(ws) / (fid + ".result.json")
    proc = run_cli([str(FLOW_CLI), str(spec_path), "-o", str(out), "--figure-number",
                    str(number), "--preview", "--json", str(res_path)], timeout=300)
    res = _load(res_path, {}) or {}
    res["ok"] = proc.returncode == 0 and out.is_file()
    res["log"] = (proc.stdout[-3000:] + proc.stderr[-1000:]).strip()
    if res["ok"]:
        (flow_out(ws) / (fid + ".svg")).write_text(svg_of(out), encoding="utf-8")
        res["svg"] = "/api/flow-preview/%s.svg" % fid
        res["png"] = "/api/flow-preview/%s.png" % fid
    _dump(res_path, res)
    return res


def delete_flowchart(ws, fid: str) -> None:
    for p in [flow_dir(ws) / (fid + ".json")] + list(flow_out(ws).glob(fid + ".*")):
        if p.is_file():
            p.unlink()


def next_flow_id(ws) -> str:
    n = 1
    while (flow_dir(ws) / ("flow%d.json" % n)).exists():
        n += 1
    return "flow%d" % n


# --------------------------------------------------------------------------- #
# AI 起草术语                                                                    #
# --------------------------------------------------------------------------- #

def glossary_of(ws) -> Dict[str, str]:
    data = _load(ws.root / "glossary.json", {}) or {}
    return {str(k): str(v) for k, v in (data.get("terms") or data).items()
            if isinstance(v, str)} if isinstance(data, dict) else {}


def bom_names_of(ws) -> Dict[str, str]:
    data = _load(ws.root / "bom.json", {}) or {}
    return {k: v.get("name", "") for k, v in (data.get("matched") or {}).items()}


def apply_terms(ws, rows: List[dict], overwrite: bool = False) -> dict:
    """把 [{selector, term, label}] 写进 plan.terms：默认只填空白，人填过的不动。

    WP1 R1 #3：读-改-写走 project_io.update_plan，在该 plan 文件的进程级锁内完成并原子落盘，
    不再与别的写者交错覆盖（原先直接 _load → 修改 → _dump，两次并发 apply 会丢一次）。"""
    from patent_figure import project_io as _pio
    outcome = {"filled": [], "kept_human": [], "added": []}

    def mutate(plan):
        by_sel = {t.get("selector"): t for t in plan.get("terms", [])}
        for r in rows:
            sel, term = r.get("selector"), str(r.get("term", "")).strip()
            if not sel or not term:
                continue
            row = by_sel.get(sel)
            if row is None:
                row = {"selector": sel, "term": term, "label": r.get("label", "once")}
                plan.setdefault("terms", []).append(row)
                by_sel[sel] = row
                outcome["added"].append(sel)
                continue
            if (row.get("term") or "").strip() and not overwrite:
                outcome["kept_human"].append(sel)
                continue
            row["term"] = term
            if r.get("label") in ("once", "all", "none"):
                row["label"] = r["label"]
            outcome["filled"].append(sel)
        return plan

    plan = _pio.update_plan(ws.plan, mutate, snapshot_cb=ws.snapshot_plan)
    return {"plan": plan, "filled": outcome["filled"],
            "kept_human": outcome["kept_human"], "added": outcome["added"]}


# --------------------------------------------------------------------------- #
# 导出：说明书附图 Word + 附图说明                                                #
# --------------------------------------------------------------------------- #

def export_bundle(ws, result: dict, dest: Path, run_cli: Callable, want_dwg: bool,
                  png_dpi: int = EXPORT_PNG_DPI) -> dict:
    files, notes = [], []
    figs = sorted(result.get("figures", []), key=lambda f: f.get("number", 999))
    # R1c #1: the browser export reads the pinned current version directory via current.json,
    # not the mutable out view, so the whole package comes from one consistent version.
    from patent_figure import versions as _VS
    src_dir = _VS.VersionStore(ws.root).current_dir() or ws.out
    # WP1-5: raster PNGs are generated here, at export, by dpi — the interactive path only
    # ever made SVGs. Render each annotated/filing DXF to PNG in-process (matplotlib is
    # imported lazily inside render_preview); a failure is cosmetic and must not block export.
    try:
        from patent_figure import sheet as _sheet
        for fig in figs:
            for suffix in ("_annotated", "_filing"):
                dxf = src_dir / (fig["id"] + suffix + ".dxf")
                png = src_dir / (fig["id"] + suffix + ".png")
                if dxf.is_file() and not png.is_file():
                    try:
                        _sheet.render_preview(dxf, png, dpi=png_dpi)
                    except Exception as exc:
                        notes.append("PNG 未生成 %s：%s" % (png.name, exc))
    except Exception as exc:  # pragma: no cover - import guard
        notes.append("预览渲染不可用：%s" % exc)
    for fig in figs:
        for suffix in ("", "_annotated", "_filing"):
            for ext in (".dxf", ".png", ".svg"):
                src = src_dir / (fig["id"] + suffix + ext)
                if src.is_file():
                    sub = {"": "原始出图", "_annotated": "交底版_带件号表",
                           "_filing": "递交版_仅图号"}[suffix]
                    (dest / sub).mkdir(parents=True, exist_ok=True)
                    shutil.copy2(src, dest / sub / src.name)
                    files.append("%s/%s" % (sub, src.name))
    flows = [f for f in list_flowcharts(ws) if f.get("result", {}) and f["result"].get("ok")]
    for f in flows:
        for ext in (".dxf", ".png", ".svg"):
            src = flow_out(ws) / (f["id"] + ext)
            if src.is_file():
                (dest / "流程图").mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, dest / "流程图" / src.name)
                files.append("流程图/" + src.name)
    numerals = src_dir / "reference-numerals.json"
    lines = []
    if numerals.is_file():
        shutil.copy2(numerals, dest / numerals.name)
        files.append(numerals.name)
        lines.append(_load(numerals, {}).get("description_zh", ""))
    desc = [f.get("annotation", {}).get("figure_description", "") for f in figs]
    desc += [f["result"].get("figure_description", "") for f in flows]
    desc = [d for d in desc if d]
    if desc:
        desc[-1] = desc[-1].rstrip("；") + "。"
    text = "【附图说明】\n" + "\n".join(desc) + "\n\n" + "\n".join(lines) + "\n"
    (dest / "附图说明.txt").write_text(text, encoding="utf-8")
    files.append("附图说明.txt")
    # 说明书附图.docx：每图一页（交底版 PNG + 流程图 PNG），A4 版心等比缩放
    pngs = [src_dir / (f["id"] + "_annotated.png") for f in figs]
    pngs += [flow_out(ws) / (f["id"] + ".png") for f in flows]
    pngs = [p for p in pngs if p.is_file()]
    if pngs:
        try:
            _docx(pngs, dest / "说明书附图_交底版.docx")
            files.append("说明书附图_交底版.docx")
        except Exception as exc:  # python-docx 不在时只少一个文件
            notes.append("Word 未生成：%s" % exc)
    dwg_log = []
    if want_dwg:
        for sub in ("交底版_带件号表", "递交版_仅图号"):
            for dxf in sorted((dest / sub).glob("*.dxf")):
                dwg = dxf.with_suffix(".dwg")
                proc = run_cli([str(SCRIPTS / "autocad_core_dxf_to_dwg.py"), str(dxf), str(dwg)],
                               timeout=600)
                if proc.returncode != 0:
                    proc = run_cli([str(SCRIPTS / "libredwg_dxf_to_dwg.py"), str(dxf),
                                    "-o", str(dxf.parent)], timeout=600)
                    dwg_log.append("%s/%s: AutoCAD 失败，改用 LibreDWG（exit=%d）"
                                   % (sub, dxf.stem, proc.returncode))
                else:
                    dwg_log.append("%s/%s: AutoCAD 转换成功" % (sub, dxf.stem))
                if dwg.is_file():
                    files.append("%s/%s" % (sub, dwg.name))
    return {"files": sorted(set(files)), "dwg_log": dwg_log, "notes": notes,
            "figure_descriptions": desc}


def _docx(pngs: List[Path], path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    from PIL import Image
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.left_margin = Cm(2.5), Cm(2.5)
    sec.bottom_margin, sec.right_margin = Cm(1.5), Cm(1.5)
    max_w, max_h = 17.0, 25.0
    for i, png in enumerate(pngs):
        with Image.open(png) as im:
            w, h = im.size
        width = min(max_w, max_h * w / h)
        para = doc.paragraphs[0] if i == 0 and doc.paragraphs else doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(png), width=Cm(width))
        if i < len(pngs) - 1:
            doc.add_page_break()
    doc.save(str(path))


# --------------------------------------------------------------------------- #
# 结构识别：给人看的结构组 + 零件中文显示名（只存 structure.json，不进 plan）       #
# --------------------------------------------------------------------------- #

def structure_file(ws) -> Path:
    return ws.root / "structure.json"


def load_structure(ws) -> dict:
    """AI 结果优先；没有就按装配层级兜底分组。新出现的零件补进「其他零件」。"""
    import workbench_llm
    asm = _load(ws.assembly, {}) or {}
    plan = _load(ws.plan, {}) or {}
    data = _load(structure_file(ws), None)
    if not (isinstance(data, dict) and data.get("groups")):
        return workbench_llm.fallback_structure(asm, plan)
    skip = workbench_llm.excluded_by_plan(plan)
    known = {p["name"] for p in asm.get("parts", [])
             if not p.get("degenerate") and not skip(p["name"])}
    placed = {n for g in data["groups"] for n in g.get("parts", [])}
    missing = sorted(known - placed)
    if missing:
        other = next((g for g in data["groups"] if g.get("name") == "其他零件"), None)
        if other is None:
            other = {"id": "g%d" % (len(data["groups"]) + 1), "name": "其他零件",
                     "role": "识别之后新出现的零件", "parts": []}
            data["groups"].append(other)
        other["parts"] += missing
    for g in data["groups"]:
        g["parts"] = [n for n in g.get("parts", []) if n in known]
    data["groups"] = [g for g in data["groups"] if g["parts"]]
    return data


def rename_groups(ws, edits: List[dict]) -> dict:
    data = load_structure(ws)
    by_id = {g["id"]: g for g in data["groups"]}
    for e in edits:
        g = by_id.get(e.get("id"))
        if not g:
            continue
        if str(e.get("name", "")).strip():
            g["name"] = str(e["name"]).strip()[:20]
        if "role" in e:
            g["role"] = str(e.get("role") or "").strip()[:60]
    _dump(structure_file(ws), data)
    return data


def name_rows(ws) -> List[dict]:
    data = load_structure(ws)
    return [{"selector": n, "term": v.get("name", ""), "label": "once"}
            for n, v in (data.get("names") or {}).items() if v.get("name")]
